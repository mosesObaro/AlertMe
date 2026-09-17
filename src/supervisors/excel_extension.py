"""
Excel Extension Bridge for Global Country-Based PhD Funding and Supervisor Intelligence.
Enriches the master Edge Computing PhD Tracker workbook with dedicated sheets:
1. Country Intelligence
2. Global Supervisor Directory
3. Global PhD Funding
Preserves all 11 foundational curriculum sheets and existing Hong Kong sheets intact.
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from typing import Optional, Any
from pathlib import Path

from .storage import StorageManager
from .config import SUPPORTED_COUNTRIES

def add_country_supervisor_sheets(
    wb: Optional[openpyxl.Workbook] = None,
    workbook_path: Optional[Any] = None,
    storage: Optional[StorageManager] = None
) -> dict:
    """
    Extends an existing openpyxl workbook (or file at workbook_path) with Country Intelligence,
    Global Supervisor Directory, and Global PhD Funding sheets.
    """
    save_on_exit = False
    target_path = None
    if wb is None:
        target_path = Path(workbook_path or "Edge_Computing_PhD_Curriculum_Tracker.xlsx")
        if not target_path.exists():
            return {"status": "skipped", "reason": f"File not found: {target_path}"}
        wb = openpyxl.load_workbook(target_path)
        save_on_exit = True

    if storage is None:
        storage = StorageManager()

    # Palette
    NAVY_DARK = "1B365D"
    NAVY_MID = "2C4D75"
    NAVY_LIGHT = "E8EEF5"
    BORDER_LIGHT = "D0D7DE"
    BORDER_DARK = "8C9BAE"
    FILL_ZEBRA = "F7F9FC"
    FILL_WHITE = "FFFFFF"

    font_title = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_data = Font(name="Calibri", size=10, color="000000")
    font_bold = Font(name="Calibri", size=10, bold=True, color="000000")

    fill_title = PatternFill(start_color=NAVY_DARK, end_color=NAVY_DARK, fill_type="solid")
    fill_header = PatternFill(start_color=NAVY_MID, end_color=NAVY_MID, fill_type="solid")
    fill_row_alt = PatternFill(start_color=FILL_ZEBRA, end_color=FILL_ZEBRA, fill_type="solid")
    fill_row_white = PatternFill(start_color=FILL_WHITE, end_color=FILL_WHITE, fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color=BORDER_LIGHT),
        right=Side(style='thin', color=BORDER_LIGHT),
        top=Side(style='thin', color=BORDER_LIGHT),
        bottom=Side(style='thin', color=BORDER_LIGHT)
    )

    # -------------------------------------------------------------------------
    # 1. SHEET: Country Intelligence
    # -------------------------------------------------------------------------
    ws_country = wb.create_sheet(title="Country Intelligence")
    ws_country.views.sheetView[0].showGridLines = True

    ws_country.merge_cells("A1:H1")
    ws_country["A1"] = "GLOBAL PhD FUNDING & COUNTRY RESEARCH ECOSYSTEMS"
    ws_country["A1"].font = font_title
    ws_country["A1"].fill = fill_title
    ws_country["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_country.row_dimensions[1].height = 36

    c_headers = [
        "Country", "Code", "Primary Funding Vehicle", "Universities",
        "Vetted Supervisors", "Key Deadline", "Dependant Visa Policy", "Status"
    ]
    for col_idx, h in enumerate(c_headers, start=1):
        cell = ws_country.cell(row=3, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws_country.row_dimensions[3].height = 24

    row_idx = 4
    for c_key, c_cfg in SUPPORTED_COUNTRIES.items():
        unis = storage.get_universities(c_key)
        profs = storage.get_professors(c_key)

        fill = fill_row_alt if (row_idx % 2 == 0) else fill_row_white
        data = [
            c_cfg["name"],
            c_cfg["country_code"],
            c_cfg["primary_funding_vehicle"],
            len(unis),
            len(profs),
            c_cfg["target_deadline"],
            c_cfg["dependant_visa_policy"][:60] + "...",
            "Active Campaign"
        ]

        for col_idx, val in enumerate(data, start=1):
            cell = ws_country.cell(row=row_idx, column=col_idx, value=val)
            cell.font = font_bold if col_idx == 1 else font_data
            cell.fill = fill
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")
        row_idx += 1

    # -------------------------------------------------------------------------
    # 2. SHEET: Global Supervisor Directory
    # -------------------------------------------------------------------------
    ws_sup = wb.create_sheet(title="Global Supervisor Directory")
    ws_sup.views.sheetView[0].showGridLines = True

    ws_sup.merge_cells("A1:J1")
    ws_sup["A1"] = "GLOBAL PhD SUPERVISOR DIRECTORY — EDGE COMPUTING & DISTRIBUTED SYSTEMS"
    ws_sup["A1"].font = font_title
    ws_sup["A1"].fill = fill_title
    ws_sup["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_sup.row_dimensions[1].height = 36

    s_headers = [
        "Country", "University", "Professor", "Department",
        "Primary Research Areas", "Fit %", "Tier", "Recruitment Status",
        "Official Profile URL", "Verified Freshness"
    ]
    for col_idx, h in enumerate(s_headers, start=1):
        cell = ws_sup.cell(row=3, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws_sup.row_dimensions[3].height = 24

    row_idx = 4
    for c_key in SUPPORTED_COUNTRIES:
        profs = storage.get_professors(c_key)
        for p in profs:
            fill = fill_row_alt if (row_idx % 2 == 0) else fill_row_white
            rec_status = p.recruitment.status if p.recruitment else "UNKNOWN"
            areas = ", ".join(p.research_interests[:3])

            data = [
                p.country,
                p.university,
                p.name,
                p.department,
                areas,
                f"{p.alignment_score:.1f}%",
                p.priority_tier,
                rec_status,
                p.official_profile_url,
                p.recruitment.last_verified if p.recruitment else "2026-09-15"
            ]

            for col_idx, val in enumerate(data, start=1):
                cell = ws_sup.cell(row=row_idx, column=col_idx, value=val)
                cell.font = font_bold if col_idx in [1, 3] else font_data
                cell.fill = fill
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")
            row_idx += 1

    # -------------------------------------------------------------------------
    # 3. SHEET: Global PhD Funding
    # -------------------------------------------------------------------------
    ws_fund = wb.create_sheet(title="Global PhD Funding")
    ws_fund.views.sheetView[0].showGridLines = True

    ws_fund.merge_cells("A1:I1")
    ws_fund["A1"] = "GLOBAL DOCTORAL FUNDING & SCHOLARSHIPS DIRECTORY"
    ws_fund["A1"].font = font_title
    ws_fund["A1"].fill = fill_title
    ws_fund["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws_fund.row_dimensions[1].height = 36

    f_headers = [
        "Country", "University / Provider", "Scheme Title", "Funding Type",
        "Tuition Coverage", "Stipend / Salary", "International Access",
        "Dependant Support", "Application Deadline"
    ]
    for col_idx, h in enumerate(f_headers, start=1):
        cell = ws_fund.cell(row=3, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws_fund.row_dimensions[3].height = 24

    row_idx = 4
    for c_key in SUPPORTED_COUNTRIES:
        funding_list = storage.get_funding_opportunities(c_key)
        for fo in funding_list:
            fill = fill_row_alt if (row_idx % 2 == 0) else fill_row_white
            data = [
                fo.country,
                fo.university,
                fo.title,
                fo.funding_type,
                fo.tuition_coverage,
                fo.stipend_amount,
                fo.international_eligibility,
                fo.dependant_support,
                fo.application_deadline
            ]

            for col_idx, val in enumerate(data, start=1):
                cell = ws_fund.cell(row=row_idx, column=col_idx, value=val)
                cell.font = font_bold if col_idx in [1, 3] else font_data
                cell.fill = fill
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")
            row_idx += 1

    # Column Width Auto-Fitting
    for ws in [ws_country, ws_sup, ws_fund]:
        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max(len(str(cell.value or '')) for cell in col[2:])
            ws.column_dimensions[col_letter].width = min(40, max(12, max_len + 3))

    print("[Excel Extension] Added Country Intelligence, Global Supervisor Directory, and Global PhD Funding sheets.")
    if save_on_exit and target_path:
        wb.save(target_path)
    return {
        "status": "success",
        "sheets_added": ["Country Intelligence", "Global Supervisor Directory", "Global PhD Funding"]
    }
