"""
Integration test verifying full 14-sheet workbook generation with Hong Kong PhD module enabled.
"""

import unittest
import tempfile
import os
import openpyxl
import sys

# Ensure scratch and scripts directories are accessible
PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
for p in [os.path.join(PROJECT_ROOT, "scratch"), os.path.join(PROJECT_ROOT, "scripts")]:
    if p not in sys.path:
        sys.path.insert(0, p)

from build_edge_curriculum import create_edge_computing_workbook

class TestHKExcelIntegration(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.output_xlsx = os.path.join(self.temp_dir, "full_14_sheets.xlsx")

    def test_full_14_sheets_generated(self):
        """Verifies all 14 sheets are created when HK module is enabled."""
        create_edge_computing_workbook(output_filename=self.output_xlsx, enable_hk_module=True)
        self.assertTrue(os.path.exists(self.output_xlsx))

        wb = openpyxl.load_workbook(self.output_xlsx)
        self.assertEqual(len(wb.sheetnames), 14)

        # 3 new additive sheets
        self.assertIn("HK Supervisor Intelligence", wb.sheetnames)
        self.assertIn("HK Scholarships & Deadlines", wb.sheetnames)
        self.assertIn("HK Application Tracker", wb.sheetnames)

        # HK Supervisor sheet has data
        ws_sup = wb["HK Supervisor Intelligence"]
        self.assertGreaterEqual(ws_sup.max_row, 12)
        self.assertEqual(ws_sup.cell(row=3, column=1).value, "University")

        # Dashboard has Hong Kong section header
        dash = wb["Dashboard"]
        self.assertIn("HONG KONG PhD CAMPAIGN", str(dash["A30"].value))

if __name__ == "__main__":
    unittest.main()
