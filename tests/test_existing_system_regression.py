"""
Regression test suite proving that the existing Edge Computing PhD Curriculum system
remains 100% functional, intact, and unaffected when the Hong Kong module is disabled.
"""

import unittest
import tempfile
import os
import openpyxl
import sys

# Ensure scratch and scripts directories are accessible
PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
for p in [os.path.join(PROJECT_ROOT, "scratch"), os.path.join(PROJECT_ROOT, "scripts")]:
    if p not in sys.path:
        sys.path.insert(0, p)

from build_edge_curriculum import create_edge_computing_workbook

class TestExistingSystemRegression(unittest.TestCase):

    def setUp(self):
        self.temp_dir = tempfile.mkdtemp()
        self.output_xlsx = os.path.join(self.temp_dir, "regression_curriculum.xlsx")

    def test_original_11_sheets_intact_when_hk_module_disabled(self):
        """Proves that disabling the Hong Kong module generates strictly the original 11 sheets."""
        create_edge_computing_workbook(output_filename=self.output_xlsx, enable_hk_module=False)
        self.assertTrue(os.path.exists(self.output_xlsx))

        wb = openpyxl.load_workbook(self.output_xlsx)
        expected_original_sheets = [
            'Dashboard',
            'Master Curriculum',
            '24-Week Plan',
            'Weekly Progress Tracker',
            'Course Tracker',
            'Paper Tracker',
            'Research Ideas',
            'Skills Matrix',
            'Simulation & Tools Guide',
            'Recommended Path',
            'PhD Readiness Assessment'
        ]

        self.assertEqual(len(wb.sheetnames), 11)
        self.assertEqual(wb.sheetnames, expected_original_sheets)

        # Assert specific original formulas exist
        dash = wb['Dashboard']
        self.assertIn("AVERAGE", str(dash['A6'].value))
        self.assertIn("COUNTIF", str(dash['G6'].value))
        self.assertIn("PhD Readiness Assessment", str(dash['K6'].value))

        # Assert original 24-Week plan has 27 rows (headers + 24 weeks)
        plan = wb['24-Week Plan']
        self.assertEqual(plan.max_row, 27)

        # Assert no Hong Kong sheets exist when disabled
        for sheet_name in wb.sheetnames:
            self.assertFalse(sheet_name.startswith("HK "))

if __name__ == "__main__":
    unittest.main()
